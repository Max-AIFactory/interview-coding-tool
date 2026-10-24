"""Exports into exports/<timestamp>/.

Every function code is rendered with its full name.  Confidential ranges are
never part of any export (segments cannot overlap them; paragraph dumps are
masked).  Every exported text is checked against the anonymization map and
remaining hits are reported as warnings (and written to WARNINGS.md)."""
from __future__ import annotations

import csv
import datetime as _dt
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..db import Database
from . import analysis
from .anonymization import list_mapping, check_text
from .core import (EN_LABEL, NQ_LABEL, list_interactions, code_list, list_memos, memo_versions,
                   changelog as _changelog)

EXPORT_TYPES = ["codebook", "interview_overview", "segments", "quote_bank", "interactions", "matrices",
                "changelog", "reliability", "memos", "full_dump"]


def _label(code: dict[str, Any]) -> str:
    return code["name"]


def _md_escape(s: Any) -> str:
    return str(s if s is not None else "").replace("|", "\\|").replace("\n", " ")


def _md_table(headers: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for r in rows:
        out.append("| " + " | ".join(_md_escape(c) for c in r) + " |")
    return "\n".join(out) + "\n"


def _write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow(["" if c is None else c for c in r])


def _write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[list[Any]]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        for i, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None][:200])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
        ws.freeze_panes = "A2"
    wb.save(path)


def _cite(s: dict[str, Any]) -> str:
    return f"{s['interview_id']}, ¶{s['para_no']}"


def _valence_label(v: int) -> str:
    return {1: "+1", -1: "-1", 0: "0"}.get(v, str(v))


class Exporter:
    def __init__(self, db: Database, exports_dir: Path):
        self.db = db
        self.exports_dir = Path(exports_dir)
        self.mapping = list_mapping(db)
        self.warnings: list[str] = []
        self.files: list[str] = []
        self.codes = code_list(db, include_retired=True)
        self.code_by_id = {c["id"]: c for c in self.codes}
        self.tree_codes = analysis._codes_in_order(db, include_retired=True)

    # -- helpers ---------------------------------------------------------
    def _check(self, text: str, where: str) -> None:
        for term in check_text(text or "", self.mapping):
            msg = f"Mapped term '{term}' still occurs in {where}."
            if msg not in self.warnings:
                self.warnings.append(msg)

    def _save(self, name: str, content: str) -> None:
        p = self.folder / name
        p.write_text(content, encoding="utf-8")
        self.files.append(name)
        self._check(content, name)

    def _segments(self) -> list[dict[str, Any]]:
        return analysis.retrieve(self.db)

    # -- individual exports ---------------------------------------------
    def codebook(self) -> None:
        headers = ["Code", "Full name", "Type", "Parent", "Definition", "Coding rule", "Anchor example",
                   "Frequency total", "Frequency DE", "Frequency CN", "Status"]
        rows = []
        missing = []
        for c in self.tree_codes:
            parent = self.code_by_id.get(c["parent_id"])
            f = c["frequency"]
            rows.append([c["code_key"], c["name"], c["type"], parent["name"] if parent else "",
                         c["definition"], c["coding_rule"], c["anchor_example"], f["total"], f["DE"], f["CN"],
                         "active" if c["active"] else "retired"])
            if c["active"] and (not c["anchor_example"].strip() or not c["definition"].strip()):
                missing.append(c["name"])
        md = "# Category system (codebook)\n\n"
        md += f"Exported {self.stamp}. Frequencies = number of coded segments in the primary coding pass.\n\n"
        if missing:
            md += "> **Incomplete codes (missing definition or anchor example):** " + ", ".join(missing) + "\n\n"
        md += _md_table(headers, rows)
        self._save("codebook.md", md)
        _write_xlsx(self.folder / "codebook.xlsx", [("Codebook", headers, rows)])
        self.files.append("codebook.xlsx")

    def interview_overview(self) -> None:
        headers = ["ID", "Country", "Role category", "Role label", "Date", "Duration (min)", "Original language",
                   "Mode", "Pilot", "Status", "Coded segments"]
        rows = []
        for i in self.db.query("SELECT * FROM interviews ORDER BY country, id"):
            n = self.db.scalar("""SELECT COUNT(*) FROM segments s JOIN paragraphs p ON p.id=s.paragraph_id
                                  JOIN coding_passes cp ON cp.id=s.coding_pass_id
                                  WHERE p.interview_id=? AND cp.name='primary'""", (i["id"],))
            rows.append([i["id"], i["country"], i["role_category"], i["role_label"], i["date"], i["duration_min"],
                         i["language_original"], i["interview_mode"], "yes" if i["pilot"] else "no", i["status"], n])
        self._save("interview_overview.md", "# Interview overview\n\n" + _md_table(headers, rows))
        _write_xlsx(self.folder / "interview_overview.xlsx", [("Interviews", headers, rows)])
        self.files.append("interview_overview.xlsx")

    def segments(self) -> None:
        segs = self._segments()
        headers = ["segment_ref", "Segment ID", "Interview", "Paragraph", "Timestamp", "Start", "End", "Locator",
                   "Country", "Role label", "Code", "Code key", "Code type", "Valence", "Agency", "Quote candidate",
                   "Quotation block", "Original English", "Comment", "Text"]
        rows = []
        for s in segs:
            for c in s["codes"]:
                rows.append([s["segment_ref"], s["id"], s["interview_id"], s["para_no"], s["timestamp"] or "",
                             s["start_char"], s["end_char"], _cite(s), s["country"], s["role_label"],
                             c["name"], c["code_key"], c["type"], _valence_label(c["valence"]), c["agency"] or "",
                             "yes" if s["quote_candidate"] else "no", NQ_LABEL[s["quotation_block"]],
                             EN_LABEL[s["en_original"]], s["comment"], s["text_snapshot"]])
                self._check(s["text_snapshot"] + " " + s["comment"], f"segment {s['id']} ({_cite(s)})")
        _write_csv(self.folder / "segments.csv", headers, rows)
        self.files.append("segments.csv")
        # by function
        by_code: dict[int, dict[str, list]] = defaultdict(lambda: {"DE": [], "CN": []})
        for s in segs:
            for c in s["codes"]:
                by_code[c["code_id"]][s["country"]].append((s, c))
        md = f"# Coded segments per category and country\n\nExported {self.stamp}. Primary coding pass only.\n\n"
        for code in self.tree_codes:
            groups = by_code.get(code["id"])
            if not groups:
                continue
            md += f"## {code['name']}\n\n"
            for cc in ("DE", "CN"):
                items = groups[cc]
                md += f"### {cc} ({len(items)} segment{'s' if len(items) != 1 else ''})\n\n"
                for s, c in items:
                    meta = f"valence {_valence_label(c['valence'])}" + (f", agency {c['agency']}" if c["agency"] else "")
                    if s["en_original"] != "none":
                        meta += "; original English" if s["en_original"] == "full" else "; partly original English"
                    if s["quotation_block"] != "none":
                        meta += "; QUOTATION BLOCK" if s["quotation_block"] == "full" else "; partly quotation-blocked"
                    ts = f", {s['timestamp']}" if s["timestamp"] else ""
                    md += (f"- **{_cite(s)}**{ts} ({s['segment_ref']}; {s['role_label']}; {meta}): "
                           f"{s['text_snapshot'].replace(chr(10), ' ')}")
                    if s["comment"]:
                        md += f"  \n  _Comment: {s['comment']}_"
                    md += "\n"
                md += "\n"
        self._save("segments_by_function.md", md)

    def quote_bank(self, segment_ids: list[int] | None = None) -> None:
        segs = [s for s in self._segments() if s["quote_candidate"]]
        if segment_ids:
            wanted = set(segment_ids)
            segs = [s for s in analysis.retrieve(self.db) if s["id"] in wanted]
        # quotation-blocked passages ([NO-QUOTE]) are never quoted verbatim
        blocked = [s for s in segs if s["quotation_block"] != "none"]
        if blocked:
            self.warnings.append(
                f"{len(blocked)} segment(s) excluded from the quote bank because they overlap a quotation-blocked "
                f"passage: " + ", ".join(s["segment_ref"] for s in blocked) + ".")
        segs = [s for s in segs if s["quotation_block"] == "none"]
        headers = ["segment_ref", "Code", "Country", "Locator", "Timestamp", "Role label", "Valence", "Agency",
                   "Original English", "Quote", "Comment"]
        rows = []
        by_code: dict[int, dict[str, list]] = defaultdict(lambda: {"DE": [], "CN": []})
        for s in segs:
            for c in s["codes"]:
                by_code[c["code_id"]][s["country"]].append((s, c))
                rows.append([s["segment_ref"], c["name"], s["country"], _cite(s), s["timestamp"] or "",
                             s["role_label"], _valence_label(c["valence"]), c["agency"] or "",
                             EN_LABEL[s["en_original"]], s["text_snapshot"], s["comment"]])
        md = (f"# Quote bank\n\nExported {self.stamp}. Attribution: role label + interview ID and paragraph. "
              f"Segments overlapping a quotation-blocked passage are excluded.\n\n")
        for code in self.tree_codes:
            g = by_code.get(code["id"])
            if not g:
                continue
            md += f"## {code['name']}\n\n"
            for cc in ("DE", "CN"):
                if not g[cc]:
                    continue
                md += f"### {cc}\n\n"
                for s, c in g[cc]:
                    en = {"full": "; original English", "partial": "; partly original English"}.get(s["en_original"], "")
                    ts = f", {s['timestamp']}" if s["timestamp"] else ""
                    md += (f"> {s['text_snapshot'].replace(chr(10), ' ')}\n>\n"
                           f"> — {s['role_label']} ({_cite(s)}{ts}; {s['segment_ref']}{en})\n\n")
        self._save("quote_bank.md", md)
        _write_csv(self.folder / "quote_bank.csv", headers, rows)
        self.files.append("quote_bank.csv")

    def interactions(self) -> None:
        view = analysis.interactions_view(self.db)
        headers = ["ID", "Country", "Interview", "From", "To", "Mechanism", "Description", "Condition", "Segment locator",
                   "Timestamp", "Segment text"]
        rows = []
        for it in view["interactions"]:
            loc = f"{it['interview_id']}, ¶{it['para_no']}" if it.get("para_no") else (it["interview_id"] or "")
            rows.append([it["id"], it["country"], it["interview_id"], it["from_name"], it["to_name"], it["mechanism"],
                         it["description"], it["condition"], loc, it.get("timestamp") or "",
                         it.get("segment_text") or ""])
        _write_csv(self.folder / "interactions.csv", headers, rows)
        self.files.append("interactions.csv")
        md = f"# Function interactions\n\nExported {self.stamp}.\n\n" + _md_table(headers, rows)
        for cc in ("DE", "CN"):
            md += f"\n## Aggregated matrix {cc} (rows = from, columns = to; inducing / blocking)\n\n"
            fk = [f["code_key"] for f in view["functions"]]
            names = {f["code_key"]: f["name"] for f in view["functions"]}
            md += _md_table(["from \\ to"] + [names[k] for k in fk],
                            [[names[a]] + [f"{view['matrix'][cc][a][b]['inducing']} / {view['matrix'][cc][a][b]['blocking']}"
                                           for b in fk] for a in fk])
        self._save("interactions.md", md)
        sheets = []
        fk = [f["code_key"] for f in view["functions"]]
        names = {f["code_key"]: f["name"] for f in view["functions"]}
        for cc in ("DE", "CN"):
            for mech in ("inducing", "blocking"):
                sheets.append((f"{cc} {mech}", ["from \\ to"] + [names[k] for k in fk],
                               [[names[a]] + [view["matrix"][cc][a][b][mech] for b in fk] for a in fk]))
        sheets.append(("Interactions", headers, rows))
        _write_xlsx(self.folder / "interactions_matrix.xlsx", sheets)
        self.files.append("interactions_matrix.xlsx")
        self._save("interactions.mmd", view["mermaid"] + "\n")

    def matrices(self) -> None:
        ag = analysis.agency_matrix(self.db)
        headers = ["Function", "Country"] + ag["agencies"]
        rows = []
        for r in ag["rows"]:
            for cc in ("DE", "CN"):
                rows.append([r["name"], cc] + [r["cells"][cc][a] for a in ag["agencies"]])
        _write_xlsx(self.folder / "agency_matrix.xlsx", [("Agency", headers, rows)])
        self.files.append("agency_matrix.xlsx")
        for attr in ("country", "role_category"):
            ct = analysis.crosstab(self.db, attr)
            headers = ["Code", "Code key"]
            for v in ct["values"]:
                headers += [f"{v} segments", f"{v} interviews", f"{v} +1", f"{v} 0", f"{v} -1"]
            rows = []
            for r in ct["rows"]:
                row = [r["name"], r["code_key"]]
                for v in ct["values"]:
                    c = r["cells"][v]
                    row += [c["segments"], c["interviews"], c["valence"]["1"], c["valence"]["0"], c["valence"]["-1"]]
                rows.append(row)
            _write_csv(self.folder / f"crosstab_{attr}.csv", headers, rows)
            self.files.append(f"crosstab_{attr}.csv")
        co = analysis.cooccurrence(self.db)
        names = [c["name"] for c in co["codes"]]
        ids = [c["id"] for c in co["codes"]]
        rows = []
        for a in ids:
            row = [self.code_by_id[a]["name"]]
            for b in ids:
                if a == b:
                    row.append(co["diagonal"].get(str(a), 0))
                else:
                    k = f"{min(a, b)}:{max(a, b)}"
                    row.append(co["cells"].get(k, 0))
            rows.append(row)
        _write_csv(self.folder / "cooccurrence.csv", ["Code"] + names, rows)
        self.files.append("cooccurrence.csv")
        m = analysis.matrix(self.db)
        headers = ["Code", "Code key"] + [i["id"] for i in m["interviews"]] + ["Total"]
        rows = []
        for c in m["codes"]:
            cells = m["cells"].get(str(c["id"]), {})
            rows.append([c["name"], c["code_key"]] + [cells.get(i["id"], 0) for i in m["interviews"]]
                        + [m["totals"].get(str(c["id"]), 0)])
        _write_csv(self.folder / "code_matrix.csv", headers, rows)
        self.files.append("code_matrix.csv")

    def changelog(self) -> None:
        rows = _changelog(self.db)
        md = f"# Codebook changelog\n\nExported {self.stamp}.\n\n"
        md += _md_table(["Date", "Phase", "Code", "Action", "Field", "Old value", "New value", "Reason"],
                        [[r["timestamp"], r["phase"], r["code_name"] or "", r["action"], r["field"] or "",
                          r["old_value"] or "", r["new_value"] or "", r["reason"]] for r in reversed(rows)])
        self._save("codebook_changelog.md", md)

    def reliability(self) -> None:
        md = f"# Intra-coder reliability report\n\nExported {self.stamp}.\n\n"
        any_report = False
        for iv in self.db.query("SELECT id FROM interviews ORDER BY country, id"):
            try:
                rep = analysis.reliability_report(self.db, iv["id"])
            except Exception:
                continue
            any_report = True
            md += f"## {iv['id']} — {rep['pass1']['name']} vs. {rep['pass2']['name']}\n\n"
            md += (f"Respondent paragraphs: {rep['paragraphs']} · identical code sets: {rep['agreeing_paragraphs']} "
                   f"({rep['percent_agreement']} %) · pooled Cohen's kappa: {rep['pooled_kappa']}\n\n")
            md += _md_table(["Code", "Both", "Only pass 1", "Only pass 2", "Neither", "% agreement", "Kappa"],
                            [[r["name"], r["both"], r["only_pass1"], r["only_pass2"], r["neither"],
                              r["percent_agreement"], r["kappa"]] for r in rep["per_code"]])
            if rep["disagreements"]:
                md += "\n### Disagreements\n\n"
                md += _md_table(["Paragraph", "Pass 1", "Pass 2", "Resolution note", "Text"],
                                [[f"¶{d['para_no']}", ", ".join(d["pass1"]), ", ".join(d["pass2"]), d["note"],
                                  d["text"][:300]] for d in rep["disagreements"]])
            md += "\n"
        if not any_report:
            md += "_No re-coding pass has been completed yet._\n"
        self._save("reliability_report.md", md)

    def memos(self) -> None:
        md = f"# Memos\n\nExported {self.stamp}.\n\n"
        for m in list_memos(self.db):
            ref = ""
            if m["scope"] == "code":
                c = self.code_by_id.get(int(m["ref_id"])) if str(m["ref_id"]).isdigit() else None
                ref = f"code: {c['name']}" if c else f"code {m['ref_id']}"
            elif m["scope"] == "interview":
                ref = f"interview: {m['ref_id']}"
            elif m["scope"] == "segment":
                ref = f"segment: {m['ref_id']}"
            else:
                ref = "general"
            md += f"## Memo {m['id']} ({ref})\n\n{m['text']}\n\n"
            vs = memo_versions(self.db, m["id"])
            if len(vs) > 1:
                md += "<details><summary>Version history</summary>\n\n"
                for v in vs[1:]:
                    md += f"**{v['saved_at']}**\n\n{v['text']}\n\n"
                md += "</details>\n\n"
        self._save("memos.md", md)

    def full_dump(self) -> None:
        conf = analysis.confidential_by_paragraph(self.db)
        dump: dict[str, Any] = {"exported_at": self.stamp, "schema_version": self.db.schema_version()}
        dump["interviews"] = self.db.query("SELECT * FROM interviews ORDER BY id")
        pars = self.db.query("SELECT * FROM paragraphs ORDER BY interview_id, idx")
        for p in pars:
            p["text"] = analysis.mask_confidential(p["text"], conf.get(p["id"], []))
            p["confidential_ranges"] = len(conf.get(p["id"], []))
        dump["paragraphs"] = pars  # includes the paragraph 'timestamp' column
        dump["en_original_ranges"] = self.db.query("SELECT * FROM en_original_ranges ORDER BY id")
        dump["no_quote_ranges"] = self.db.query("SELECT * FROM no_quote_ranges ORDER BY id")
        dump["codes"] = self.db.query("SELECT * FROM codes ORDER BY sort_order, id")
        dump["coding_passes"] = self.db.query("SELECT * FROM coding_passes ORDER BY id")
        dump["segments"] = self.db.query(
            "SELECT s.*, (p.interview_id || '-S' || s.id) AS segment_ref, p.timestamp "
            "FROM segments s JOIN paragraphs p ON p.id = s.paragraph_id ORDER BY s.id")
        dump["segment_codes"] = self.db.query("SELECT * FROM segment_codes ORDER BY segment_id, code_id")
        dump["interactions"] = self.db.query("SELECT * FROM interactions ORDER BY id")
        dump["memos"] = self.db.query("SELECT * FROM memos ORDER BY id")
        dump["memo_versions"] = self.db.query("SELECT * FROM memo_versions ORDER BY id")
        dump["code_changelog"] = self.db.query("SELECT * FROM code_changelog ORDER BY id")
        dump["reliability_notes"] = self.db.query("SELECT * FROM reliability_notes ORDER BY id")
        dump["settings"] = {r["key"]: r["value"] for r in self.db.query("SELECT key, value FROM settings")
                            if r["key"] not in ("secondary_backup_dir",)}
        text = json.dumps(dump, ensure_ascii=False, indent=1)
        self._save("full_dump.json", text)
        d = self.folder / "full_dump_csv"
        d.mkdir(exist_ok=True)
        for name in ("interviews", "paragraphs", "en_original_ranges", "no_quote_ranges", "codes", "coding_passes",
                     "segments", "segment_codes", "interactions", "memos", "memo_versions", "code_changelog",
                     "reliability_notes"):
            rows = dump[name]
            headers = list(rows[0].keys()) if rows else []
            _write_csv(d / f"{name}.csv", headers, [[r[h] for h in headers] for r in rows])
            self.files.append(f"full_dump_csv/{name}.csv")

    # -- driver ----------------------------------------------------------
    def run(self, types: list[str] | None = None, quote_segment_ids: list[int] | None = None) -> dict[str, Any]:
        types = types or EXPORT_TYPES
        bad = [t for t in types if t not in EXPORT_TYPES]
        if bad:
            raise ValueError(f"Unknown export type(s): {', '.join(bad)}. Allowed: {', '.join(EXPORT_TYPES)}")
        now = _dt.datetime.now()
        self.stamp = now.replace(microsecond=0).isoformat(sep=" ")
        self.folder = self.exports_dir / now.strftime("%Y-%m-%d_%H%M%S")
        n = 0
        while self.folder.exists():
            n += 1
            self.folder = self.exports_dir / (now.strftime("%Y-%m-%d_%H%M%S") + f"_{n}")
        self.folder.mkdir(parents=True)
        for t in types:
            if t == "quote_bank":
                self.quote_bank(quote_segment_ids)
            else:
                getattr(self, t)()
        if self.warnings:
            (self.folder / "WARNINGS.md").write_text(
                "# Export warnings\n\n" + "\n".join(f"- {w}" for w in self.warnings) + "\n", encoding="utf-8")
            self.files.append("WARNINGS.md")
        if set(types) >= set(EXPORT_TYPES):
            self.db.set_setting("last_full_export", self.stamp)
        return {"path": str(self.folder), "files": sorted(self.files), "warnings": self.warnings, "types": types}


def run_export(db: Database, exports_dir: Path, types: list[str] | None = None,
               quote_segment_ids: list[int] | None = None) -> dict[str, Any]:
    return Exporter(db, exports_dir).run(types, quote_segment_ids)
