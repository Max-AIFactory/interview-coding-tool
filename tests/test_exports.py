"""Exports never contain confidential text or mapped terms; F-codes rendered with full
names; quote export filtered by country contains role labels (sections 8, 9.6, 9.7)."""
import csv
import io
import json
from pathlib import Path

from conftest import code_quote, locate

CONF_DE = "legal dispute with one specific grid operator"
CONF_CN = "target for next year has not been published"


def _all_text(folder: Path) -> str:
    out = []
    for f in folder.rglob("*"):
        if f.is_file() and f.suffix in (".md", ".csv", ".json", ".mmd"):
            out.append(f.read_text(encoding="utf-8"))
        elif f.suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(f, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
    return "\n".join(out)


def test_confidential_never_exported_and_not_in_matrices(demo):
    """Acceptance test 7."""
    client, codes = demo
    # confidential text cannot be coded, so matrices cannot contain it; also mark a manual range
    hit = locate(client, "CN-00", "Inland provinces started only last year.")
    r = client.post("/api/v1/confidential", json={"paragraph_id": hit["paragraph_id"], "start_char": hit["start_char"], "end_char": hit["end_char"]})
    assert r.status_code == 201
    code_quote(client, codes, "CN-00", "Jiangsu and Guangdong have several years of experience", ["F4"])
    code_quote(client, codes, "DE-00", "Overall the direction is positive.", ["F5"], valence=1, quote_candidate=True)
    r = client.post("/api/v1/export", json={})
    assert r.status_code == 200, r.text
    folder = Path(r.json()["path"])
    text = _all_text(folder)
    assert CONF_DE not in text and CONF_CN not in text and "Inland provinces started only last year" not in text
    dump = json.loads((folder / "full_dump.json").read_text(encoding="utf-8"))
    assert any("[CONFIDENTIAL]" in p["text"] for p in dump["paragraphs"])
    # search does not surface confidential text either
    assert client.get("/api/v1/search?q=legal dispute").json() == []
    assert client.get("/api/v1/search?q=Inland provinces").json() == []
    # matrices ignore confidential text (no segment can exist there)
    m = client.get("/api/v1/analysis/matrix").json()
    assert sum(m["totals"].values()) == 2


def test_quote_export_cn_legitimation_has_role_labels_and_no_real_names(demo):
    """Acceptance test 6."""
    client, codes = demo
    client.post("/api/v1/anonymization", json={"term": "Jiangsu", "replacement": "[Province A]"})
    code_quote(client, codes, "CN-00", "Trust has to be earned event by event.", ["F5"], valence=0, quote_candidate=True)
    code_quote(client, codes, "DE-00", "Today at least the transmission operators treat us as partners", ["F5"], valence=1, quote_candidate=True)
    # a quote that still contains a mapped term -> export must warn
    s = code_quote(client, codes, "CN-00", "Jiangsu and Guangdong have several years of experience", ["F5"], quote_candidate=True)
    segs = client.get(f"/api/v1/retrieve?code_id={codes['F5']['id']}&country=CN&quote_candidate=true").json()
    assert {x["id"] for x in segs} == {s["id"]} | {x["id"] for x in segs if x["id"] != s["id"]}
    r = client.post("/api/v1/export", json={"types": ["quote_bank"], "quote_segment_ids": [x["id"] for x in segs]})
    assert r.status_code == 200
    folder = Path(r.json()["path"])
    md = (folder / "quote_bank.md").read_text(encoding="utf-8")
    assert "Manager at a provincial grid company" in md and "CN-00, ¶" in md
    assert "Trust has to be earned" in md and "transmission operators" not in md      # filtered to CN
    assert "## Legitimation" in md and "F5" not in md.replace("CN-00", "")          # full names, no bare code
    assert any("Jiangsu" in w for w in r.json()["warnings"])
    assert (folder / "WARNINGS.md").exists()
    # after applying the anonymization the warning disappears
    client.post("/api/v1/anonymization/apply", json={})
    r = client.post("/api/v1/export", json={"types": ["quote_bank", "segments"]})
    assert r.json()["warnings"] == []
    assert "[Province A]" in (Path(r.json()["path"]) / "segments.csv").read_text(encoding="utf-8-sig")


def test_anonymization_apply_shifts_offsets(demo):
    client, codes = demo
    client.post("/api/v1/anonymization", json={"term": "Bundesnetzagentur", "replacement": "[the regulator]"})
    s = code_quote(client, codes, "DE-00", "Overall the direction is positive.", ["F5"])      # after the term
    s2 = code_quote(client, codes, "DE-00", "The regulator, the Bundesnetzagentur, listens", ["F5"])  # contains the term
    hits = client.post("/api/v1/anonymization/scan").json()
    assert len(hits) == 1 and hits[0]["term"] == "Bundesnetzagentur"
    r = client.post("/api/v1/anonymization/apply", json={"hits": hits})
    assert r.json()["replaced"] == 1
    g = client.get(f"/api/v1/segments/{s['id']}").json()
    assert g["text_snapshot"] == "Overall the direction is positive."
    g2 = client.get(f"/api/v1/segments/{s2['id']}").json()
    assert g2["text_snapshot"] == "The regulator, the [the regulator], listens"
    iv = client.get("/api/v1/interviews/DE-00").json()
    par = next(p for p in iv["paragraphs"] if p["id"] == s["paragraph_id"])
    assert par["text"][g["start_char"]:g["end_char"]] == g["text_snapshot"]
    assert par["confidential"] and par["text"][par["confidential"][0]["start_char"]:par["confidential"][0]["end_char"]].startswith("We currently have a legal dispute")


def test_function_codes_rendered_with_full_names_everywhere(demo):
    client, codes = demo
    code_quote(client, codes, "DE-00", "The technical side works.", ["F1"], valence=1)
    code_quote(client, codes, "DE-00", "Capital was never the main issue", ["F6"])
    client.post("/api/v1/interactions", json={"from_code_id": codes["F4"]["id"], "to_code_id": codes["F6"]["id"], "mechanism": "blocking", "country": "DE", "interview_id": "DE-00"})
    r = client.post("/api/v1/export", json={})
    folder = Path(r.json()["path"])
    for name in ("segments_by_function.md", "interactions.md", "codebook.md", "interactions.mmd"):
        t = (folder / name).read_text(encoding="utf-8")
        assert "Knowledge Development" in t or "Market Formation" in t
    seg_md = (folder / "segments_by_function.md").read_text(encoding="utf-8")
    assert "## Knowledge Development" in seg_md and "## Resource Mobilization" in seg_md
    assert "## F1" not in seg_md and "## F6" not in seg_md
    mmd = (folder / "interactions.mmd").read_text(encoding="utf-8")
    assert '"Market Formation"' in mmd and "blocking (1)" in mmd
    # API views carry names too
    for kind in ("agency", "interactions", "barriers"):
        j = client.get(f"/api/v1/analysis/{kind}").json()
        assert "Knowledge Development" in json.dumps(j)
    cb = (folder / "codebook.md").read_text(encoding="utf-8")
    assert "| F1 | Knowledge Development |" in cb and "Incomplete codes" in cb    # anchors missing in seed


def test_en_original_reported_on_segments_and_exports(demo):
    client, codes = demo
    iv = client.get("/api/v1/interviews/DE-00").json()
    en_pars = [p for p in iv["paragraphs"] if p["en_original"]]
    assert len(en_pars) == 1
    e = en_pars[0]["en_original"][0]
    assert en_pars[0]["text"][e["start_char"]:e["end_char"]] == "That speed is something we could learn from."
    full = code_quote(client, codes, "DE-00", "That speed is something we could learn from.", ["ML"])
    assert full["en_original"] == "full"
    part = code_quote(client, codes, "DE-00", "within a year. That speed", ["ML"])
    assert part["en_original"] == "partial"
    none = code_quote(client, codes, "DE-00", "The technical side works.", ["F1"], quote_candidate=True)
    assert none["en_original"] == "none"
    client.patch(f"/api/v1/segments/{full['id']}", json={"quote_candidate": True})
    r = client.post("/api/v1/export", json={"types": ["segments", "quote_bank", "full_dump"]})
    assert r.status_code == 200, r.text
    folder = Path(r.json()["path"])
    csv_text = (folder / "segments.csv").read_text(encoding="utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(csv_text)))
    by_id = {int(x["Segment ID"]): x["Original English"] for x in rows}
    assert by_id[full["id"]] == "yes" and by_id[part["id"]] == "partly" and by_id[none["id"]] == "no"
    qb = (folder / "quote_bank.md").read_text(encoding="utf-8")
    assert "; original English)" in qb
    qcsv = list(csv.DictReader(io.StringIO((folder / "quote_bank.csv").read_text(encoding="utf-8-sig"))))
    assert {x["Original English"] for x in qcsv} == {"yes", "no"}
    dump = json.loads((folder / "full_dump.json").read_text(encoding="utf-8"))
    assert len(dump["en_original_ranges"]) == 1
    segs = client.get(f"/api/v1/retrieve?code_id={codes['ML']['id']}").json()
    assert {s["en_original"] for s in segs} == {"full", "partial"}
    # anonymization keeps the range aligned
    client.post("/api/v1/anonymization", json={"term": "Jiangsu", "replacement": "[Province A]"})
    client.post("/api/v1/anonymization/apply", json={})
    iv = client.get("/api/v1/interviews/DE-00").json()
    p = next(p for p in iv["paragraphs"] if p["en_original"])
    e = p["en_original"][0]
    assert p["text"][e["start_char"]:e["end_char"]] == "That speed is something we could learn from."


def test_export_file_set_and_reminder(demo):
    client, codes = demo
    assert client.get("/api/v1/dashboard").json()["weekly_reminder"] is True
    r = client.post("/api/v1/export", json={})
    files = set(r.json()["files"])
    for f in ("codebook.md", "codebook.xlsx", "interview_overview.md", "interview_overview.xlsx", "segments_by_function.md",
              "segments.csv", "quote_bank.md", "quote_bank.csv", "interactions.csv", "interactions.md", "interactions_matrix.xlsx",
              "interactions.mmd", "agency_matrix.xlsx", "crosstab_country.csv", "crosstab_role_category.csv", "cooccurrence.csv",
              "code_matrix.csv", "codebook_changelog.md", "reliability_report.md", "memos.md", "full_dump.json",
              "full_dump_csv/segments.csv"):
        assert f in files, f
    assert client.get("/api/v1/dashboard").json()["weekly_reminder"] is False
    assert client.post("/api/v1/export", json={"types": ["nope"]}).status_code == 422
