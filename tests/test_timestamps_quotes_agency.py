"""Paragraph timestamps, [NO-QUOTE] quotation block, agency value list, segment_ref, migrations."""
import csv
import io
import json
import sqlite3
from pathlib import Path

import pytest

from app.db import SCHEMA_SQL, SCHEMA_VERSION, AGENCY_VALUES, Database
from app.main import Project
from app.seed import SEED_CODES
from app.services.core import QUOTE_BLOCK_MESSAGE
from conftest import BASE, SAMPLES, code_quote, import_md, locate

NQ_FULL = "One of the big utilities tried to have us excluded from the working group last year"
NQ_PART = "quite a stir. We exchange experiences on prequalification"
CLEAN = "We exchange experiences on prequalification, and that has shortened the learning curve"


def _rows(folder: Path, name: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO((folder / name).read_text(encoding="utf-8-sig"))))


# ---------------------------------------------------------------------------
# 1. timestamps
# ---------------------------------------------------------------------------

def test_timestamps_on_paragraphs_segments_locate_and_search(demo):
    client, codes = demo
    iv = client.get("/api/v1/interviews/DE-00").json()
    assert all(p["timestamp"] for p in iv["paragraphs"])
    cn = client.get("/api/v1/interviews/CN-00").json()
    assert all(p["timestamp"] is None for p in cn["paragraphs"])          # CN-00 has no timestamps -> null
    hits = client.post("/api/v1/interviews/DE-00/locate", json={"text": "Capital was never the main issue"}).json()
    assert hits[0]["timestamp"] == "00:17:40"
    # timestamp filter: exact match on the paragraph timestamp
    assert client.post("/api/v1/interviews/DE-00/locate", json={"text": "the", "timestamp": "00:17:40"}).json()
    assert all(h["timestamp"] == "00:17:40" for h in
               client.post("/api/v1/interviews/DE-00/locate", json={"text": "the", "timestamp": "00:17:40"}).json())
    assert client.post("/api/v1/interviews/DE-00/locate", json={"text": "Capital", "timestamp": "00:00:20"}).json() == []
    s = code_quote(client, codes, "DE-00", "Capital was never the main issue", ["F6"])
    assert s["timestamp"] == "00:17:40" and s["segment_ref"] == f"DE-00-S{s['id']}"
    assert client.get(f"/api/v1/segments/{s['id']}").json()["timestamp"] == "00:17:40"
    assert client.get("/api/v1/retrieve").json()[0]["timestamp"] == "00:17:40"
    assert client.get(f"/api/v1/codes/{codes['F6']['id']}/segments").json()[0]["timestamp"] == "00:17:40"
    iv = client.get("/api/v1/interviews/DE-00").json()
    seg = [x for p in iv["paragraphs"] for x in p["segments"]][0]
    assert seg["timestamp"] == "00:17:40"
    hit = client.get("/api/v1/search?q=Capital was never").json()[0]
    assert hit["timestamp"] == "00:17:40"
    cn_seg = code_quote(client, codes, "CN-00", "We learn mostly by doing", ["F1"])
    assert cn_seg["timestamp"] is None


def test_timestamp_columns_in_exports(demo):
    client, codes = demo
    s = code_quote(client, codes, "DE-00", "The technical side works.", ["F1"], quote_candidate=True)
    client.post("/api/v1/interactions", json={"segment_id": s["id"], "from_code_id": codes["F1"]["id"],
                                              "to_code_id": codes["F4"]["id"], "mechanism": "inducing"})
    r = client.post("/api/v1/export", json={"types": ["segments", "quote_bank", "interactions", "full_dump"]})
    assert r.status_code == 200, r.text
    folder = Path(r.json()["path"])
    seg_rows = _rows(folder, "segments.csv")
    assert list(seg_rows[0].keys())[0] == "segment_ref"
    assert seg_rows[0]["segment_ref"] == f"DE-00-S{s['id']}" and seg_rows[0]["Timestamp"] == "00:02:31"
    qb = _rows(folder, "quote_bank.csv")
    assert qb[0]["segment_ref"] == f"DE-00-S{s['id']}" and qb[0]["Timestamp"] == "00:02:31"
    assert f"DE-00-S{s['id']}" in (folder / "quote_bank.md").read_text(encoding="utf-8")
    it = _rows(folder, "interactions.csv")
    assert it[0]["Timestamp"] == "00:02:31"
    dump = json.loads((folder / "full_dump.json").read_text(encoding="utf-8"))
    de_first = next(x for x in dump["paragraphs"] if x["interview_id"] == "DE-00" and x["idx"] == 0)
    assert de_first["timestamp"] == "00:00:05"
    assert dump["segments"][0]["timestamp"] == "00:02:31" and dump["segments"][0]["segment_ref"] == f"DE-00-S{s['id']}"
    assert "timestamp" in _rows(folder, "full_dump_csv/paragraphs.csv")[0]
    assert "timestamp" in _rows(folder, "full_dump_csv/segments.csv")[0]


# ---------------------------------------------------------------------------
# 2. quotation block
# ---------------------------------------------------------------------------

def test_no_quote_ranges_in_interview_view_and_derived_status(demo):
    client, codes = demo
    iv = client.get("/api/v1/interviews/DE-00").json()
    pars = [p for p in iv["paragraphs"] if p["no_quote"]]
    assert len(pars) == 1
    nq = pars[0]["no_quote"][0]
    assert pars[0]["text"][nq["start_char"]:nq["end_char"]].startswith(NQ_FULL)
    full = code_quote(client, codes, "DE-00", NQ_FULL, ["F5"], valence=-1)      # coding IS allowed
    part = code_quote(client, codes, "DE-00", NQ_PART, ["F1b"])
    none = code_quote(client, codes, "DE-00", CLEAN, ["F1b"], valence=1)
    assert (full["quotation_block"], part["quotation_block"], none["quotation_block"]) == ("full", "partial", "none")
    assert client.get(f"/api/v1/segments/{full['id']}").json()["quotation_block"] == "full"
    by_id = {s["id"]: s["quotation_block"] for s in client.get("/api/v1/retrieve").json()}
    assert by_id[full["id"]] == "full" and by_id[part["id"]] == "partial" and by_id[none["id"]] == "none"
    by_id = {s["id"]: s["quotation_block"] for s in client.get(f"/api/v1/codes/{codes['F1b']['id']}/segments").json()}
    assert by_id == {part["id"]: "partial", none["id"]: "none"}
    iv = client.get("/api/v1/interviews/DE-00").json()
    view = {s["id"]: s["quotation_block"] for p in iv["paragraphs"] for s in p["segments"]}
    assert view[full["id"]] == "full"


def test_quote_candidate_refused_on_blocked_segments(demo):
    client, codes = demo
    loc = locate(client, "DE-00", NQ_FULL)
    body = {"paragraph_id": loc["paragraph_id"], "start_char": loc["start_char"], "end_char": loc["end_char"],
            "codes": [{"code_id": codes["F5"]["id"]}], "quote_candidate": True}
    r = client.post("/api/v1/segments", json=body)
    assert r.status_code == 422 and r.json()["detail"] == QUOTE_BLOCK_MESSAGE
    assert client.get("/api/v1/interviews/DE-00").json()["segment_count"] == 0
    body["quote_candidate"] = False
    seg = client.post("/api/v1/segments", json=body).json()
    r = client.patch(f"/api/v1/segments/{seg['id']}", json={"quote_candidate": True})
    assert r.status_code == 422 and r.json()["detail"] == QUOTE_BLOCK_MESSAGE
    assert client.get(f"/api/v1/segments/{seg['id']}").json()["quote_candidate"] is False
    # partial overlap counts too
    loc2 = locate(client, "DE-00", NQ_PART)
    r = client.post("/api/v1/segments", json={**body, "start_char": loc2["start_char"], "end_char": loc2["end_char"],
                                              "quote_candidate": True})
    assert r.status_code == 422 and "quotation-blocked" in r.json()["detail"]
    # a clean quote candidate cannot be moved into the blocked range
    clean = code_quote(client, codes, "DE-00", CLEAN, ["F1b"], quote_candidate=True)
    r = client.patch(f"/api/v1/segments/{clean['id']}", json={"start_char": loc2["start_char"], "end_char": clean["end_char"]})
    assert r.status_code == 422 and r.json()["detail"] == QUOTE_BLOCK_MESSAGE
    g = client.get(f"/api/v1/segments/{clean['id']}").json()
    assert g["start_char"] == clean["start_char"] and g["quote_candidate"] is True
    # moving it while dropping the flag in the same request is fine; other patches on blocked segments are fine
    assert client.patch(f"/api/v1/segments/{seg['id']}", json={"comment": "(F5-) excluded from working group"}).status_code == 200
    # batch ops: create_segment and patch_segment with the flag are refused atomically
    ops = [{"op": "patch_segment", "segment_id": clean["id"], "comment": "x"},
           {"op": "create_segment", "paragraph_id": loc["paragraph_id"], "start_char": loc["start_char"],
            "end_char": loc["end_char"], "codes": [{"code_id": codes["F1b"]["id"]}], "quote_candidate": True}]
    r = client.post("/api/v1/segments/batch", json={"operations": ops})
    assert r.status_code == 422 and "Operation 1" in r.json()["detail"] and QUOTE_BLOCK_MESSAGE in r.json()["detail"]
    assert client.get(f"/api/v1/segments/{clean['id']}").json()["comment"] == ""       # rolled back
    r = client.post("/api/v1/segments/batch", json={"operations": [
        {"op": "patch_segment", "segment_id": seg["id"], "quote_candidate": True}]})
    assert r.status_code == 422 and QUOTE_BLOCK_MESSAGE in r.json()["detail"]


def test_quotation_block_in_retrieval_and_exports(demo):
    client, codes = demo
    full = code_quote(client, codes, "DE-00", NQ_FULL, ["F5"], valence=-1)
    part = code_quote(client, codes, "DE-00", NQ_PART, ["F5"])
    clean = code_quote(client, codes, "DE-00", CLEAN, ["F5"], quote_candidate=True)
    # simulate a legacy flag on a blocked segment (bypassing the API rule) -> retrieval and quote bank still exclude it
    with client.app.state.project.db.write("legacy flag", record_undo=False) as tx:
        tx.execute("UPDATE segments SET quote_candidate=1 WHERE id IN (?,?)", (full["id"], part["id"]))
    ids = {s["id"] for s in client.get("/api/v1/retrieve?quote_candidate=true").json()}
    assert ids == {clean["id"]}
    ids = {s["id"] for s in client.get(f"/api/v1/retrieve?code_id={codes['F5']['id']}&quote_candidate=true").json()}
    assert ids == {clean["id"]}
    r = client.post("/api/v1/export", json={"types": ["segments", "quote_bank", "full_dump"]})
    assert r.status_code == 200, r.text
    folder = Path(r.json()["path"])
    rows = {int(x["Segment ID"]): x for x in _rows(folder, "segments.csv")}
    assert rows[full["id"]]["Quotation block"] == "yes" and rows[part["id"]]["Quotation block"] == "partly"
    assert rows[clean["id"]]["Quotation block"] == "no"
    qb_md = (folder / "quote_bank.md").read_text(encoding="utf-8")
    assert CLEAN in qb_md and NQ_FULL not in qb_md and "quite a stir" not in qb_md
    qb = _rows(folder, "quote_bank.csv")
    assert {int(x["segment_ref"].rsplit("S", 1)[1]) for x in qb} == {clean["id"]}
    assert any("quotation-blocked" in w for w in r.json()["warnings"])
    # explicit ids from the retrieval view are filtered as well
    r = client.post("/api/v1/export", json={"types": ["quote_bank"], "quote_segment_ids": [full["id"], clean["id"]]})
    md = (folder.parent / Path(r.json()["path"]).name / "quote_bank.md").read_text(encoding="utf-8")
    assert NQ_FULL not in md and CLEAN in md
    dump = json.loads((folder / "full_dump.json").read_text(encoding="utf-8"))
    assert len(dump["no_quote_ranges"]) == 1
    assert "full_dump_csv/no_quote_ranges.csv" in r.json()["files"] or (folder / "full_dump_csv" / "no_quote_ranges.csv").exists()
    by_fn = (folder / "segments_by_function.md").read_text(encoding="utf-8")
    assert "QUOTATION BLOCK" in by_fn and "partly quotation-blocked" in by_fn


def test_anonymization_shifts_no_quote_ranges(demo):
    client, codes = demo
    client.post("/api/v1/anonymization", json={"term": "Verband", "replacement": "[the association]"})
    client.post("/api/v1/anonymization/apply", json={})
    iv = client.get("/api/v1/interviews/DE-00").json()
    p = next(p for p in iv["paragraphs"] if p["no_quote"])
    nq = p["no_quote"][0]
    assert p["text"][nq["start_char"]:nq["end_char"]].startswith(NQ_FULL) and "[the association]" in p["text"]


# ---------------------------------------------------------------------------
# 3. agency value list
# ---------------------------------------------------------------------------

def test_agency_value_list_everywhere(demo):
    client, codes = demo
    assert list(AGENCY_VALUES) == ["GOV-CENTRAL", "GOV-PROV", "REGULATOR", "GRID", "SOE", "MARKET", "ASSOC", "RESEARCH", "MULTI"]
    st = client.get("/api/v1/settings").json()
    assert st["agencies"] == list(AGENCY_VALUES) and st["agency_meanings"]["GRID"] == "grid companies, TSO, DSO"
    schema = client.get("/openapi.json").json()["components"]["schemas"]["CodeAssignment"]["properties"]["agency"]
    enums = [x["enum"] for x in schema.get("anyOf", [schema]) if "enum" in x]
    assert enums and enums[0] == list(AGENCY_VALUES)
    hit = locate(client, "DE-00", "Capital was never the main issue")
    for old in ("public", "private", "co-steering", "status quo/unclear"):
        r = client.post("/api/v1/segments", json={"paragraph_id": hit["paragraph_id"], "start_char": hit["start_char"],
                                                  "end_char": hit["end_char"], "codes": [{"code_id": codes["F6"]["id"], "agency": old}]})
        assert r.status_code == 422, old
    s = code_quote(client, codes, "DE-00", "Capital was never the main issue", ["F6"], agency="MARKET")
    assert s["codes"][0]["agency"] == "MARKET"
    assert client.patch(f"/api/v1/segments/{s['id']}/codes/{codes['F6']['id']}", json={"agency": "GOV-PROV"}).json()["codes"][0]["agency"] == "GOV-PROV"
    assert client.patch(f"/api/v1/segments/{s['id']}/codes/{codes['F6']['id']}", json={"agency": "private"}).status_code == 422
    assert client.patch(f"/api/v1/segments/{s['id']}/codes/{codes['F6']['id']}", json={"agency": ""}).json()["codes"][0]["agency"] is None
    r = client.post("/api/v1/segments/batch", json={"operations": [
        {"op": "patch_code", "segment_id": s["id"], "code_id": codes["F6"]["id"], "agency": "MULTI"}]})
    assert r.status_code == 200 and r.json()[0]["codes"][0]["agency"] == "MULTI"
    code_quote(client, codes, "CN-00", "The market is not really a market yet", ["F4"], valence=-1, agency="GOV-PROV")
    code_quote(client, codes, "CN-00", "Trust has to be earned event by event", ["F5"])
    assert {x["id"] for x in client.get("/api/v1/retrieve?agency=MULTI").json()} == {s["id"]}
    assert len(client.get("/api/v1/retrieve?agency=unset").json()) == 1
    ag = client.get("/api/v1/analysis/agency").json()
    assert ag["agencies"] == list(AGENCY_VALUES) + ["unset"]
    f6 = next(r for r in ag["rows"] if r["code_key"] == "F6")
    assert f6["cells"]["DE"]["MULTI"] == 1
    f4 = next(r for r in ag["rows"] if r["code_key"] == "F4")
    assert f4["cells"]["CN"]["GOV-PROV"] == 1
    r = client.post("/api/v1/export", json={"types": ["segments", "matrices"]})
    folder = Path(r.json()["path"])
    rows = _rows(folder, "segments.csv")
    assert {x["Agency"] for x in rows} == {"MULTI", "GOV-PROV", ""}
    from openpyxl import load_workbook
    ws = load_workbook(folder / "agency_matrix.xlsx", read_only=True).worksheets[0]
    header = [c for c in next(ws.iter_rows(values_only=True))]
    assert header == ["Function", "Country"] + list(AGENCY_VALUES) + ["unset"]


# ---------------------------------------------------------------------------
# 4. migration of an existing (schema v2) project
# ---------------------------------------------------------------------------

LEGACY_AGENCY = "agency TEXT CHECK (agency IS NULL OR agency IN ('public','private','co-steering','status quo/unclear'))"


def _legacy_v2_project(path: Path) -> None:
    """Build a project file exactly as version 0.8.0 left it (schema 2) with legacy agency values."""
    conn = sqlite3.connect(str(path))
    for stmt in SCHEMA_SQL.split(";"):
        if stmt.strip():
            conn.execute(stmt)
    conn.execute("ALTER TABLE paragraphs DROP COLUMN timestamp")
    conn.execute("DROP TABLE no_quote_ranges")
    conn.execute("DROP TABLE segment_codes")
    conn.execute(f"CREATE TABLE segment_codes (segment_id INTEGER NOT NULL REFERENCES segments(id), "
                 f"code_id INTEGER NOT NULL REFERENCES codes(id), valence INTEGER NOT NULL DEFAULT 0 "
                 f"CHECK (valence IN (-1,0,1)), {LEGACY_AGENCY}, PRIMARY KEY (segment_id, code_id))")
    conn.execute("CREATE INDEX ix_segcodes_code ON segment_codes(code_id)")
    conn.execute("INSERT INTO settings(key,value) VALUES ('schema_version','2'), ('revision','7')")
    for i, (key, name, parent, ctype, d, r, a, color) in enumerate(SEED_CODES[:8]):
        conn.execute("INSERT INTO codes(id, code_key, name, type, definition, coding_rule, anchor_example, color, sort_order, active) "
                     "VALUES (?,?,?,?,?,?,?,?,?,1)", (i + 1, key, name, ctype, d, r, a, color, i))
    conn.execute("INSERT INTO interviews(id, country, role_category, role_label, date, pilot) VALUES ('DE-99','DE','Research','x','2026-01-01',0)")
    conn.execute("INSERT INTO coding_passes(id, interview_id, name) VALUES (1,'DE-99','primary')")
    conn.execute("INSERT INTO paragraphs(id, interview_id, idx, para_no, speaker, text) VALUES (1,'DE-99',0,1,'R','Some respondent text here.')")
    for sid in range(1, 7):
        conn.execute("INSERT INTO segments(id, paragraph_id, start_char, end_char, text_snapshot, coding_pass_id) VALUES (?,1,0,4,'Some',1)", (sid,))
    for sid, agency in [(1, "public"), (2, "private"), (3, "co-steering"), (4, "status quo/unclear"), (5, None), (6, "public")]:
        conn.execute("INSERT INTO segment_codes(segment_id, code_id, valence, agency) VALUES (?,1,0,?)", (sid, agency))
    conn.execute("INSERT INTO undo_log(action_json, inverse_json, undone) VALUES ('{}','{}',0)")
    conn.commit()
    conn.close()


def test_migration_from_schema_2_remaps_agency_and_adds_columns(tmp_path, capsys):
    data = tmp_path / "data"; data.mkdir()
    backups = tmp_path / "backups"; backups.mkdir()
    _legacy_v2_project(data / "project.db")
    p = Project(BASE, data_dir=data, backups_dir=backups, exports_dir=tmp_path / "exports")
    out = capsys.readouterr().out
    assert "Schema upgrade 2 -> 5" in out and "backup written" in out
    assert "0003_paragraph_timestamp" in out and "0004_no_quote_ranges" in out
    assert "0005_agency_values (5 segment_codes row(s) remapped)" in out
    assert p.db.schema_version() == SCHEMA_VERSION == 5
    assert len(p.backups.list_backups()) >= 1                                  # backup before migration
    rows = {r["segment_id"]: r["agency"] for r in p.db.query("SELECT segment_id, agency FROM segment_codes ORDER BY segment_id")}
    assert rows == {1: "GOV-CENTRAL", 2: "MARKET", 3: "MULTI", 4: None, 5: None, 6: "GOV-CENTRAL"}
    assert p.db.scalar("SELECT COUNT(*) FROM undo_log") == 0                    # stale journal cleared
    assert p.db.scalar("SELECT timestamp FROM paragraphs WHERE id=1") is None    # old paragraphs keep NULL
    assert p.db.scalar("SELECT COUNT(*) FROM no_quote_ranges") == 0
    assert p.db.revision() == 7 and p.db.scalar("SELECT COUNT(*) FROM codes") == 8   # data untouched
    with pytest.raises(sqlite3.IntegrityError):
        p.db.conn.execute("INSERT INTO segment_codes(segment_id, code_id, valence, agency) VALUES (1,2,0,'public')")
    p.db.conn.execute("INSERT INTO segment_codes(segment_id, code_id, valence, agency) VALUES (1,2,0,'GRID')")
    assert p.db.integrity_check() == "ok"
    # re-opening does not migrate again
    p.db.close()
    p2 = Project(BASE, data_dir=data, backups_dir=backups, exports_dir=tmp_path / "exports")
    assert "Schema upgrade" not in capsys.readouterr().out and p2.db.schema_version() == 5
    p2.db.close()


def test_fresh_database_starts_at_current_schema(tmp_path):
    db = Database(tmp_path / "x.db")
    applied = db.migrate()
    assert db.schema_version() == SCHEMA_VERSION and len(applied) == 1 and applied[0].startswith("0001_initial")
    cols = {r[1] for r in db.conn.execute("PRAGMA table_info(paragraphs)").fetchall()}
    assert "timestamp" in cols
    assert db.conn.execute("SELECT name FROM sqlite_master WHERE name='no_quote_ranges'").fetchone()
    db.close()
